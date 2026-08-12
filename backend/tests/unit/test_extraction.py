"""Document text extraction.

Fixtures are generated in memory rather than committed as binary files. Checked-in sample
PDFs are opaque in review, can carry real personal data by accident, and make it impossible to
see from the test what shape of document is being exercised.
"""

from __future__ import annotations

import io

import pytest
from docx import Document

from app.services.extraction import (
    DOCX_CONTENT_TYPE,
    PDF_CONTENT_TYPE,
    DocumentKind,
    ExtractionError,
    detect_kind,
    extract_text,
    normalise,
)


def make_pdf(text: str) -> bytes:
    """Build a minimal single-page PDF containing `text`.

    Hand-assembled rather than pulled from a rendering library: reportlab would be a
    multi-megabyte dependency used by nothing but this helper. Byte offsets for the xref table
    are computed as objects are appended, so the file is structurally valid rather than relying
    on pypdf's tolerance for broken cross-reference tables.
    """
    escaped = text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
    stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET"

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream".encode(),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]

    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for index, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{index} 0 obj\n".encode() + body + b"\nendobj\n"

    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n".encode()
    out += f"startxref\n{xref_at}\n%%EOF\n".encode()
    return bytes(out)


def make_docx(paragraphs: list[str], table_rows: list[list[str]] | None = None) -> bytes:
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)

    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for row_index, row in enumerate(table_rows):
            for cell_index, value in enumerate(row):
                table.rows[row_index].cells[cell_index].text = value

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class TestDetectKind:
    def test_identifies_pdf_by_magic_bytes(self) -> None:
        assert detect_kind(make_pdf("hello")) is DocumentKind.PDF

    def test_identifies_docx_by_magic_bytes(self) -> None:
        assert detect_kind(make_docx(["hello"])) is DocumentKind.DOCX

    def test_rejects_a_file_lying_about_its_type(self) -> None:
        # The upload can claim application/pdf all it likes; the bytes decide. This is what
        # stops an executable reaching a parser that was never designed to see one.
        with pytest.raises(ExtractionError):
            detect_kind(b"MZ\x90\x00This is a Windows executable")

    @pytest.mark.parametrize("data", [b"", b"plain text", b"{}", b"\x00\x01\x02\x03"])
    def test_rejects_unrecognised_bytes(self, data: bytes) -> None:
        with pytest.raises(ExtractionError):
            detect_kind(data)


class TestNormalise:
    def test_collapses_runs_of_spaces(self) -> None:
        assert normalise("Python     Docker") == "Python Docker"

    def test_preserves_paragraph_breaks(self) -> None:
        assert normalise("Experience\n\n\n\n\nSkills") == "Experience\n\nSkills"

    def test_converts_windows_line_endings(self) -> None:
        assert "\r" not in normalise("line one\r\nline two")

    def test_replaces_non_breaking_spaces(self) -> None:
        # Endemic in PDF text layers, and invisible in a diff. Left alone they break naive
        # word matching, so "Node.js" fails to match a skill vocabulary entry for no visible
        # reason.
        assert normalise("Node\xa0js") == "Node js"

    def test_strips_leading_and_trailing_whitespace(self) -> None:
        assert normalise("\n\n  Ada  \n\n") == "Ada"


class TestExtractPdf:
    def test_extracts_text(self) -> None:
        kind, text = extract_text(make_pdf("Python and PostgreSQL"))

        assert kind is DocumentKind.PDF
        assert "Python" in text
        assert "PostgreSQL" in text

    def test_rejects_a_corrupt_pdf(self) -> None:
        # Right magic bytes, garbage body — the realistic shape of a truncated upload.
        with pytest.raises(ExtractionError, match="could not be read"):
            extract_text(b"%PDF-1.4\nnot actually a pdf at all")

    def test_rejects_a_pdf_with_no_text_layer(self) -> None:
        # A scanned resume is images. Reporting success with empty text would produce an empty
        # skill profile and no explanation of why.
        with pytest.raises(ExtractionError, match="No text could be extracted"):
            extract_text(make_pdf(""))


class TestExtractDocx:
    def test_extracts_paragraphs(self) -> None:
        kind, text = extract_text(make_docx(["Ada Lovelace", "Backend Engineer"]))

        assert kind is DocumentKind.DOCX
        assert "Ada Lovelace" in text
        assert "Backend Engineer" in text

    def test_extracts_table_cells(self) -> None:
        # Skills and dates very often live in tables, which `document.paragraphs` excludes.
        # Missing them would silently halve the useful text on a large minority of resumes.
        data = make_docx(["Skills"], table_rows=[["Python", "Docker"], ["SQL", "Redis"]])
        _kind, text = extract_text(data)

        for skill in ("Python", "Docker", "SQL", "Redis"):
            assert skill in text

    def test_rejects_a_corrupt_docx(self) -> None:
        with pytest.raises(ExtractionError):
            extract_text(b"PK\x03\x04 but not really a zip archive")


class TestExtractText:
    def test_rejects_an_empty_file(self) -> None:
        with pytest.raises(ExtractionError, match="empty"):
            extract_text(b"")

    def test_error_messages_are_safe_to_show_a_user(self) -> None:
        # No traceback, no file path, no library name — an error surfaced verbatim in the UI
        # must not describe the server's internals.
        with pytest.raises(ExtractionError) as caught:
            extract_text(b"%PDF-1.4\ncorrupt")

        message = str(caught.value)
        assert "Traceback" not in message
        assert "/app" not in message
        assert "pypdf" not in message.lower()


class TestContentTypeConstants:
    def test_match_the_official_media_types(self) -> None:
        assert PDF_CONTENT_TYPE == "application/pdf"
        assert DOCX_CONTENT_TYPE == (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
